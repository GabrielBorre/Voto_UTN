"""Corredor Django que exporta los resultados QA a una matriz Excel."""

import inspect
import re
import unittest
import unicodedata
from collections.abc import Mapping
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ImproperlyConfigured
from django.test.runner import DiscoverRunner
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config.pruebas_qa_relevantes import (
    CANTIDAD_PRUEBAS_RELEVANTES,
    PRUEBAS_RELEVANTES,
)

COLUMNAS_REPORTE = (
    "Componente Funcional",
    "Caso de Uso",
    "# Caso de Prueba",
    "Regresión (SI/NO)",
    "Descripción",
    "Precondiciones",
    "Datos de Entrada",
    "Pasos",
    "Resultado Esperado",
    "Diseñador",
    "Fecha",
    "Resultado",
    "Severidad",
    "Prioridad",
    "Observaciones del Resultado",
    "Tester",
)

_ALIAS_CAMPOS = {
    "componente funcional": "Componente Funcional",
    "caso de uso": "Caso de Uso",
    "caso de prueba": "# Caso de Prueba",
    "regresion": "Regresión (SI/NO)",
    "regresion si no": "Regresión (SI/NO)",
    "descripcion": "Descripción",
    "precondiciones": "Precondiciones",
    "datos de entrada": "Datos de Entrada",
    "pasos": "Pasos",
    "resultado esperado": "Resultado Esperado",
    "disenador": "Diseñador",
    "severidad": "Severidad",
    "prioridad": "Prioridad",
    "tester": "Tester",
}

_ANCHOS_COLUMNAS = (24, 25, 48, 18, 40, 34, 32, 42, 42, 22, 16, 16, 16, 16, 72, 22)


def _normalizar_etiqueta(etiqueta):
    sin_tildes = "".join(
        caracter
        for caracter in unicodedata.normalize("NFKD", etiqueta)
        if not unicodedata.combining(caracter)
    )
    sin_puntuacion = re.sub(r"[^a-zA-Z0-9]+", " ", sin_tildes)
    return re.sub(r"\s+", " ", sin_puntuacion).strip().lower()


def _normalizar_metadatos(metadatos):
    if not isinstance(metadatos, Mapping):
        raise TypeError("qa_metadata debe ser un diccionario de metadatos.")

    normalizados = {}
    for etiqueta, valor in metadatos.items():
        campo = _ALIAS_CAMPOS.get(_normalizar_etiqueta(str(etiqueta)))
        if campo is None:
            raise ValueError(f"Campo de metadatos QA no reconocido: {etiqueta!r}.")
        normalizados[campo] = valor
    return normalizados


def _metadatos_docstring(docstring):
    metadatos = {}
    campo_actual = None

    for linea in inspect.cleandoc(docstring or "").splitlines():
        contenido = linea.strip()
        if not contenido:
            campo_actual = None
            continue

        if ":" in contenido:
            etiqueta, valor = contenido.split(":", 1)
            campo = _ALIAS_CAMPOS.get(_normalizar_etiqueta(etiqueta))
            if campo is not None:
                metadatos[campo] = valor.strip()
                campo_actual = campo
                continue

        if campo_actual is not None:
            anterior = str(metadatos[campo_actual]).rstrip()
            metadatos[campo_actual] = f"{anterior}\n{contenido}".strip()

    return metadatos


def extraer_metadatos_prueba(prueba):
    """Combina metadatos de clase, método y docstring, en ese orden."""
    metadatos = _normalizar_metadatos(
        getattr(prueba.__class__, "qa_metadata", {})
    )

    nombre_metodo = getattr(prueba, "_testMethodName", "")
    metodo = getattr(prueba, nombre_metodo, None)
    metadatos.update(
        _normalizar_metadatos(getattr(metodo, "qa_metadata", {}))
    )
    metadatos.update(
        _metadatos_docstring(inspect.getdoc(metodo) if metodo else "")
    )
    return metadatos


def crear_libro_reporte(registros):
    """Crea el libro de matriz QA usando exclusivamente las 16 columnas acordadas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Matriz de pruebas"
    hoja.append(COLUMNAS_REPORTE)

    for registro in registros:
        valores = []
        for columna in COLUMNAS_REPORTE:
            valor = registro.get(columna, "")
            if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
                valor = f"'{valor}"
            valores.append(valor)
        hoja.append(valores)

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(vertical="center", wrap_text=True)

    for indice, ancho in enumerate(_ANCHOS_COLUMNAS, start=1):
        letra_columna = hoja.cell(row=1, column=indice).column_letter
        hoja.column_dimensions[letra_columna].width = ancho

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        fila[10].number_format = "dd/mm/yyyy"

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    return libro


class ResultadoPruebasExcel(unittest.TextTestResult):
    """Acumula un registro por test con el resultado y el detalle de su ejecución."""

    def __init__(self, *args, fecha_ejecucion=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.fecha_ejecucion = fecha_ejecucion or timezone.localtime()
        self.registros_reporte = []
        self._registros_por_id = {}

    def _obtener_registro(self, prueba):
        identificador = prueba.id()
        registro = self._registros_por_id.get(identificador)
        if registro is None:
            metadatos = extraer_metadatos_prueba(prueba)
            registro = {columna: "" for columna in COLUMNAS_REPORTE}
            registro.update(metadatos)
            registro["# Caso de Prueba"] = (
                registro["# Caso de Prueba"] or identificador
            )
            registro["Fecha"] = self.fecha_ejecucion.date()
            registro["Resultado"] = ""
            registro["Observaciones del Resultado"] = ""
            self._registros_por_id[identificador] = registro
            self.registros_reporte.append(registro)
        return registro

    def startTest(self, test):
        self._obtener_registro(test)
        super().startTest(test)

    def _registrar_resultado(self, prueba, resultado, detalle=""):
        registro = self._obtener_registro(prueba)
        prioridad = {"PASSED": 0, "SKIPPED": 1, "FAILED": 2, "ERROR": 3}
        resultado_actual = registro["Resultado"]
        if prioridad.get(resultado, 0) >= prioridad.get(resultado_actual, -1):
            registro["Resultado"] = resultado

        if detalle:
            observaciones = registro["Observaciones del Resultado"]
            registro["Observaciones del Resultado"] = "\n\n".join(
                parte for parte in (observaciones, detalle) if parte
            )

    def addSuccess(self, test):
        self._registrar_resultado(test, "PASSED")
        super().addSuccess(test)

    def addFailure(self, test, err):
        self._registrar_resultado(
            test, "FAILED", self._exc_info_to_string(err, test)
        )
        super().addFailure(test, err)

    def addError(self, test, err):
        self._registrar_resultado(test, "ERROR", self._exc_info_to_string(err, test))
        super().addError(test, err)

    def addSkip(self, test, reason):
        self._registrar_resultado(test, "SKIPPED", str(reason))
        super().addSkip(test, reason)

    def addExpectedFailure(self, test, err):
        self._registrar_resultado(
            test,
            "FAILED",
            f"Fallo esperado por unittest:\n{self._exc_info_to_string(err, test)}",
        )
        super().addExpectedFailure(test, err)

    def addUnexpectedSuccess(self, test):
        self._registrar_resultado(
            test,
            "FAILED",
            "La prueba marcada como fallo esperado finalizó inesperadamente.",
        )
        super().addUnexpectedSuccess(test)

    def addSubTest(self, test, subtest, err):
        if err is not None:
            es_fallo = issubclass(err[0], test.failureException)
            resultado = "FAILED" if es_fallo else "ERROR"
            self._registrar_resultado(
                test,
                resultado,
                self._exc_info_to_string(err, subtest),
            )
        super().addSubTest(test, subtest, err)


class CorredorPruebasExcel(DiscoverRunner):
    """Ejecuta la selección QA y guarda una matriz Excel al finalizar."""

    def __init__(self, *args, **kwargs):
        self._resultado_pruebas = None
        self._fecha_ejecucion = None
        super().__init__(*args, **kwargs)

    def build_suite(self, test_labels=None, **kwargs):
        suite_completa = super().build_suite(test_labels, **kwargs)
        pruebas_descubiertas = list(_iterar_pruebas(suite_completa))
        ids_relevantes = _ids_pruebas_relevantes()
        pruebas_seleccionadas = [
            prueba
            for prueba in pruebas_descubiertas
            if prueba.id() in ids_relevantes
        ]

        if not test_labels:
            ids_descubiertos = {prueba.id() for prueba in pruebas_descubiertas}
            faltantes = sorted(ids_relevantes - ids_descubiertos)
            if faltantes:
                raise ImproperlyConfigured(
                    "No se encontraron todas las pruebas QA seleccionadas. "
                    f"Faltan: {', '.join(faltantes)}"
                )
            if len(pruebas_seleccionadas) != CANTIDAD_PRUEBAS_RELEVANTES:
                raise ImproperlyConfigured(
                    "La selección QA debe contener exactamente "
                    f"{CANTIDAD_PRUEBAS_RELEVANTES} pruebas."
                )
        elif not pruebas_seleccionadas:
            raise ImproperlyConfigured(
                "Las etiquetas indicadas no contienen ninguna de las "
                "67 pruebas QA seleccionadas."
            )

        self.log(
            f"Se seleccionaron {len(pruebas_seleccionadas)} pruebas QA."
        )
        return self.test_suite(pruebas_seleccionadas)

    def get_resultclass(self):
        fecha_ejecucion = self._fecha_ejecucion

        class ResultadoConFecha(ResultadoPruebasExcel):
            def __init__(self, *args, **kwargs):
                kwargs["fecha_ejecucion"] = fecha_ejecucion
                super().__init__(*args, **kwargs)

        return ResultadoConFecha

    def run_suite(self, *args, **kwargs):
        self._resultado_pruebas = super().run_suite(*args, **kwargs)
        return self._resultado_pruebas

    def run_tests(self, test_labels=None, **kwargs):
        self._fecha_ejecucion = timezone.localtime()
        cantidad_fallos = super().run_tests(test_labels, **kwargs)
        if self._resultado_pruebas is None:
            raise RuntimeError(
                "Django no devolvió resultados; no se pudo generar el reporte Excel."
            )

        ruta_reporte = _ruta_reporte()
        ruta_reporte.parent.mkdir(parents=True, exist_ok=True)
        libro = crear_libro_reporte(self._resultado_pruebas.registros_reporte)
        libro.save(ruta_reporte)

        if self.verbosity > 0:
            print(f"\nReporte Excel de pruebas guardado en: {ruta_reporte}")
        return cantidad_fallos


def _iterar_pruebas(suite):
    for prueba in suite:
        if isinstance(prueba, unittest.TestSuite):
            yield from _iterar_pruebas(prueba)
        else:
            yield prueba


def _ids_pruebas_relevantes():
    return {
        f"VotoUTN.{clase}.{metodo}"
        for clase, metodos in PRUEBAS_RELEVANTES.items()
        for metodo in metodos
    }


def _ruta_reporte():
    ruta_configurada = getattr(settings, "TEST_REPORT_XLSX_PATH", None)
    if ruta_configurada:
        ruta = Path(ruta_configurada)
        if not ruta.is_absolute():
            ruta = Path(settings.BASE_DIR) / ruta
        return ruta

    marca_tiempo = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    return Path(settings.BASE_DIR) / f"reporte_pruebas_{marca_tiempo}.xlsx"


if CANTIDAD_PRUEBAS_RELEVANTES != 67:
    raise ImproperlyConfigured(
        "La selección de pruebas QA debe incluir exactamente 67 métodos; "
        f"actualmente incluye {CANTIDAD_PRUEBAS_RELEVANTES}."
    )
