import csv
import io
from collections import Counter
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.autoridades.models import AsignacionAutoridad, CandidaturaAutoridad
from apps.elecciones.models import Eleccion, EleccionClaustro, EleccionClaustroDepartamentoSede
from apps.mesas.models import AsignacionMesa, Mesa
from apps.padron.models import AsignacionSedePadron, RegistroPadron


PLANTILLA_AUTORIDADES_HEADERS = ("Tipo documento", "Documento", "Legajo")
PLANTILLA_AUTORIDADES_EJEMPLO = [
    ("DNI", "40123456", "2024001"),
    ("DNI", "40234567", "2024002"),
]

def leer_filas_autoridades(contenido: bytes, nombre_archivo: str = "") -> list[dict[str, str]]:
    nombre_archivo = (nombre_archivo or "").lower()
    if nombre_archivo.endswith(".xls"):
        raise ValueError("El formato .xls no es compatible. Guardá la planilla como .xlsx o CSV.")
    if nombre_archivo.endswith(".xlsx"):
        try:
            libro = load_workbook(filename=io.BytesIO(contenido), read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as exc:
            raise ValueError("No se pudo leer el archivo Excel. Verificá que sea un .xlsx válido.") from exc
        try:
            filas = list(libro.active.iter_rows(values_only=True))
            if not filas:
                return []
            encabezados = [str(valor).strip() if valor is not None else "" for valor in filas[0]]
            registros = []
            for numero, fila in enumerate(filas[1:], start=2):
                if not any(valor is not None and str(valor).strip() for valor in fila):
                    continue
                if any(
                    valor is not None
                    and str(valor).strip()
                    and (indice >= len(encabezados) or not encabezados[indice])
                    for indice, valor in enumerate(fila)
                ):
                    raise ValueError(f"La fila {numero} contiene datos sin una cabecera.")
                registros.append({
                    nombre: "" if indice >= len(fila) or fila[indice] is None else str(fila[indice]).strip()
                    for indice, nombre in enumerate(encabezados)
                })
            return registros
        finally:
            libro.close()

    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("El archivo debe estar codificado en UTF-8.") from exc

    lector = csv.DictReader(io.StringIO(texto, newline=""))
    filas = []
    for numero, fila in enumerate(lector, start=2):
        if not fila:
            continue
        if None in fila:
            raise ValueError(f"La fila {numero} contiene más columnas que las cabeceras.")
        filas.append({(clave or "").strip(): (valor or "").strip() for clave, valor in fila.items()})
    return filas


@transaction.atomic
def asignar_autoridad(registro_padron, mesa, turno, usuario):
    eleccion_claustro = mesa.eleccion_claustro_departamento.eleccion_claustro
    if not eleccion_claustro.turnos_autoridad.filter(turno=turno).exists():
        raise ValidationError("El turno de trabajo no está habilitado para este claustro.")
    configuracion_turno = eleccion_claustro.turnos_autoridad.get(turno=turno)
    if AsignacionAutoridad.objects.filter(mesa=mesa, turno=turno, tipo=AsignacionAutoridad.Tipo.TITULAR).exclude(estado=AsignacionAutoridad.Estado.RECHAZADA).count() >= configuracion_turno.titulares_por_mesa:
        raise ValidationError("La mesa ya alcanzó el máximo de autoridades configurado para este turno.")
    if not registro_padron.activo:
        raise ValidationError("Solo pueden asignarse electores activos del padrón.")
    if registro_padron.eleccion_claustro_departamento_id != mesa.eleccion_claustro_departamento_id:
        raise ValidationError("La autoridad y la mesa deben pertenecer al mismo departamento.")
    if registro_padron.sede_id and registro_padron.sede_id == mesa.sede_id:
        raise ValidationError("El elector no puede ser autoridad en la sede donde cursa.")
    if registro_padron.sede_id and not registro_padron.turno_cursada_id:
        raise ValidationError("El padrón debe tener el turno de cursada antes de asignar esta autoridad.")
    if registro_padron.turno_cursada_id and registro_padron.turno_cursada.se_superpone_con(turno):
        raise ValidationError("El turno de trabajo se superpone con el turno de cursada del elector.")
    sede_electoral = _sede_electoral_actual(registro_padron)
    if sede_electoral is not None and sede_electoral != mesa.sede_id:
        raise ValidationError("La autoridad debe asignarse en la sede electoral que ya tiene asignada.")
    candidatura, _ = CandidaturaAutoridad.objects.get_or_create(registro_padron=registro_padron, defaults={"cargada_por": usuario})
    asignacion, creada = AsignacionAutoridad.objects.get_or_create(
        registro_padron=registro_padron,
        defaults={"mesa": mesa, "turno": turno, "candidatura": candidatura, "asignada_por": usuario},
    )
    if not creada and asignacion.mesa_id != mesa.id:
        raise ValidationError("El elector ya fue asignado como autoridad de otra mesa.")
    if not creada and asignacion.turno_id != turno.id:
        raise ValidationError("El elector ya fue asignado como autoridad en otro turno.")
    asignacion.full_clean()
    if creada:
        from apps.usuarios.models import AsignacionRol, PerfilUsuario

        perfil = PerfilUsuario.objects.filter(elector=registro_padron.elector, activo=True).select_related("usuario").first()
        if perfil:
            AsignacionRol.objects.get_or_create(
                usuario=perfil.usuario,
                rol=AsignacionRol.Rol.AUTORIDAD_MESA,
                eleccion=registro_padron.eleccion,
                sede=mesa.sede,
                mesa=mesa,
            )
    return asignacion, creada


@transaction.atomic
def redistribuir_autoridades(eleccion_claustro, usuario):
    """Recalcula el plantel completo del claustro de forma determinista."""
    from apps.usuarios.models import AsignacionRol

    mesas = list(Mesa.objects.filter(
        eleccion_claustro_departamento__eleccion_claustro=eleccion_claustro,
    ).select_related("sede").order_by("numero"))
    configuraciones = list(eleccion_claustro.turnos_autoridad.select_related("turno").order_by("turno__hora_inicio", "turno__nombre"))
    if not mesas:
        raise ValidationError("Primero deben generarse las mesas del claustro.")
    if not configuraciones:
        raise ValidationError("Configurá al menos un turno de trabajo.")
    if any(config.titulares_por_mesa < 1 for config in configuraciones):
        raise ValidationError("Cada turno debe tener al menos un titular por mesa.")
    if RegistroPadron.objects.filter(eleccion=eleccion_claustro.eleccion, qr_generado_en__isnull=False).exists():
        raise ValidationError("No se pueden redistribuir autoridades después de emitir el padrón imprimible.")

    anteriores = AsignacionAutoridad.objects.filter(mesa__in=mesas)
    AsignacionRol.objects.filter(
        rol=AsignacionRol.Rol.AUTORIDAD_MESA,
        eleccion=eleccion_claustro.eleccion,
        mesa__in=mesas,
    ).delete()
    anteriores.delete()

    candidaturas = list(CandidaturaAutoridad.objects.filter(
        registro_padron__activo=True,
        registro_padron__eleccion_claustro_departamento__eleccion_claustro=eleccion_claustro,
        excluida=False,
    ).select_related(
        "registro_padron__elector", "registro_padron__sede", "registro_padron__turno_cursada",
        "registro_padron__asignacion_mesa__mesa",
    ).order_by("registro_padron__elector__apellido", "registro_padron__elector__nombre", "id"))
    plantel = [c.registro_padron for c in candidaturas]
    disponibles = {registro.pk: registro for registro in plantel}
    asignaciones = []
    sin_cubrir = []
    for config in configuraciones:
        turno = config.turno
        # Primero cubrimos la mesa de elector cuando pertenece al mismo alcance/sede.
        cubiertas = set()
        for mesa in mesas:
            candidato = _elegir_candidato(disponibles.values(), mesa, turno, solo_mesa_propia=True)
            if candidato is None:
                continue
            asignaciones.append(_crear_asignacion(candidato, mesa, turno, usuario, AsignacionAutoridad.Tipo.TITULAR))
            disponibles.pop(candidato.pk)
            cubiertas.add(mesa.pk)
        # Si la mesa propia no se puede cubrir, buscamos otra de la misma sede/departamento.
        for mesa in mesas:
            if mesa.pk in cubiertas:
                continue
            candidato = _elegir_candidato(disponibles.values(), mesa, turno)
            if candidato is None:
                sin_cubrir.append((mesa, turno))
                continue
            asignaciones.append(_crear_asignacion(candidato, mesa, turno, usuario, AsignacionAutoridad.Tipo.TITULAR))
            disponibles.pop(candidato.pk)
        # Las plazas restantes se distribuyen por rondas entre las mesas.
        for tipo, cantidad in (
            (AsignacionAutoridad.Tipo.TITULAR, max(0, config.titulares_por_mesa - 1)),
            (AsignacionAutoridad.Tipo.SUPLENTE, config.suplentes_por_mesa),
        ):
            for orden in range(cantidad):
                for mesa in mesas:
                    candidato = _elegir_candidato(disponibles.values(), mesa, turno)
                    if candidato is None:
                        continue
                    numero_suplencia = orden + 1 if tipo == AsignacionAutoridad.Tipo.SUPLENTE else None
                    asignaciones.append(_crear_asignacion(candidato, mesa, turno, usuario, tipo, numero_suplencia))
                    disponibles.pop(candidato.pk)

    AsignacionAutoridad.objects.bulk_create(asignaciones)
    nuevas = AsignacionAutoridad.objects.filter(pk__in=[a.pk for a in asignaciones]).select_related("registro_padron__elector", "mesa")
    for asignacion in nuevas:
        _sincronizar_rol_autoridad(asignacion)
    return {"asignadas": len(asignaciones), "sin_cubrir": sin_cubrir, "mesas": len(mesas)}


def _elegir_candidato(candidatos, mesa, turno, *, solo_mesa_propia=False):
    elegibles = [registro for registro in candidatos if (
        registro.activo
        and registro.eleccion_claustro_departamento_id == mesa.eleccion_claustro_departamento_id
        and _sede_electoral_actual(registro) == mesa.sede_id
        and (not registro.sede_id or registro.turno_cursada_id)
        and registro.sede_id != mesa.sede_id
        and (not registro.turno_cursada_id or not registro.turno_cursada.se_superpone_con(turno))
    )]
    if solo_mesa_propia:
        elegibles = [registro for registro in elegibles if getattr(getattr(registro, "asignacion_mesa", None), "mesa_id", None) == mesa.id]
    return min(elegibles, key=lambda registro: (
        getattr(getattr(registro, "asignacion_mesa", None), "mesa_id", None) != mesa.id,
        registro.elector.apellido.casefold(), registro.elector.nombre.casefold(), registro.elector.legajo,
    ), default=None)


def _sede_electoral_actual(registro):
    asignacion_sede = getattr(registro, "asignacion_sede_padron", None)
    if asignacion_sede:
        return asignacion_sede.sede_id
    asignacion_mesa = getattr(registro, "asignacion_mesa", None)
    return asignacion_mesa.mesa.sede_id if asignacion_mesa else None


def _crear_asignacion(registro, mesa, turno, usuario, tipo, orden=None):
    return AsignacionAutoridad(
        registro_padron=registro,
        candidatura=getattr(registro, "candidatura_autoridad", None),
        mesa=mesa,
        turno=turno,
        tipo=tipo,
        orden_suplencia=orden,
        asignada_por=usuario,
    )


def _sincronizar_rol_autoridad(asignacion):
    from apps.usuarios.models import AsignacionRol, PerfilUsuario

    perfil = PerfilUsuario.objects.filter(elector=asignacion.registro_padron.elector, activo=True).select_related("usuario").first()
    if perfil:
        AsignacionRol.objects.get_or_create(
            usuario=perfil.usuario,
            rol=AsignacionRol.Rol.AUTORIDAD_MESA,
            eleccion=asignacion.registro_padron.eleccion,
            sede=asignacion.mesa.sede,
            mesa=asignacion.mesa,
        )


def proponer_cambios_sede(eleccion_claustro, sin_cubrir):
    """Devuelve movimientos posibles sin aplicarlos, para confirmación expresa."""
    propuestas, usados = [], set()
    deficits = Counter(mesa.sede.nombre if mesa.sede_id else "" for mesa, _turno in sin_cubrir)
    ordenados = sorted(
        sin_cubrir,
        key=lambda par: (
            -deficits[par[0].sede.nombre if par[0].sede_id else ""],
            par[0].sede.nombre.casefold() if par[0].sede_id else "",
            par[0].numero,
            par[1].hora_inicio,
        ),
    )
    for mesa, turno in ordenados:
        if mesa.eleccion_claustro_departamento.eleccion_claustro_id != eleccion_claustro.id:
            continue
        candidatos = CandidaturaAutoridad.objects.filter(
            registro_padron__activo=True,
            registro_padron__eleccion_claustro_departamento=mesa.eleccion_claustro_departamento,
            registro_padron__asignacion_sede_padron__estado=AsignacionSedePadron.Estado.ASIGNADA,
            registro_padron__asignacion_mesa__isnull=False,
            excluida=False,
        ).exclude(registro_padron__asignacion_autoridad__isnull=False).exclude(
            registro_padron__asignacion_sede_padron__sede=mesa.sede,
        ).exclude(registro_padron__sede=mesa.sede).select_related(
            "registro_padron__elector", "registro_padron__sede", "registro_padron__turno_cursada",
            "registro_padron__asignacion_sede_padron__sede", "registro_padron__asignacion_mesa__mesa",
        ).order_by("registro_padron__elector__apellido", "registro_padron__elector__nombre")
        for candidatura in candidatos:
            registro = candidatura.registro_padron
            if registro.pk in usados or (registro.sede_id and not registro.turno_cursada_id) or (registro.turno_cursada_id and registro.turno_cursada.se_superpone_con(turno)):
                continue
            mesa_origen = registro.asignacion_mesa.mesa
            if mesa_origen.eleccion_claustro_departamento_id != mesa.eleccion_claustro_departamento_id:
                continue
            if AsignacionMesa.objects.filter(mesa=mesa_origen).count() <= 1:
                continue
            if not EleccionClaustroDepartamentoSede.objects.filter(
                eleccion_claustro_departamento=mesa.eleccion_claustro_departamento,
                sede=mesa.sede,
            ).exists():
                continue
            propuestas.append({
                "registro": registro,
                "candidatura": candidatura,
                "mesa_origen": mesa_origen,
                "mesa_destino": mesa,
                "turno": turno,
                "sede_origen": registro.asignacion_sede_padron.sede,
                "sede_destino": mesa.sede,
                "electores_mesa_origen": AsignacionMesa.objects.filter(mesa=mesa_origen).count(),
            })
            usados.add(registro.pk)
            break
    return propuestas


@transaction.atomic
def confirmar_cambio_sede_para_autoridad(registro, sede_destino, usuario):
    """Cambia la sede electoral y recalcula mesas y planteles dentro de una transacción."""
    from apps.usuarios.models import AsignacionRol
    from apps.padron.services import calcular_mesas_automaticas_eleccion

    eleccion = Eleccion.objects.select_for_update().get(pk=registro.eleccion_id)
    if RegistroPadron.objects.filter(eleccion=eleccion, qr_generado_en__isnull=False).exists():
        raise ValidationError("El padrón imprimible ya fue emitido; no se puede cambiar la sede electoral.")
    if not EleccionClaustroDepartamentoSede.objects.filter(
        eleccion_claustro_departamento=registro.eleccion_claustro_departamento,
        sede=sede_destino,
    ).exists():
        raise ValidationError("La sede propuesta ya no está habilitada para este departamento.")
    if registro.sede_id == sede_destino.id:
        raise ValidationError("El elector ya tiene asignada esa sede donde cursa; no puede votar allí como autoridad.")
    asignacion_sede = getattr(registro, "asignacion_sede_padron", None)
    if asignacion_sede is None:
        raise ValidationError("El elector no tiene una asignación de sede electoral que pueda modificarse.")
    asignacion_mesa = getattr(registro, "asignacion_mesa", None)
    if asignacion_mesa and AsignacionMesa.objects.filter(mesa=asignacion_mesa.mesa).count() <= 1:
        raise ValidationError("El cambio dejaría vacía su mesa electoral de origen; elegí otra persona.")

    AsignacionRol.objects.filter(rol=AsignacionRol.Rol.AUTORIDAD_MESA, eleccion=eleccion).delete()
    AsignacionAutoridad.objects.filter(registro_padron__eleccion=eleccion).delete()
    asignacion_sede.sede = sede_destino
    asignacion_sede.estado = AsignacionSedePadron.Estado.ASIGNADA
    asignacion_sede.full_clean()
    asignacion_sede.save(update_fields=("sede", "estado", "calculada_en"))
    calcular_mesas_automaticas_eleccion(eleccion, usuario=usuario)
    for claustro in EleccionClaustro.objects.filter(eleccion=eleccion):
        if CandidaturaAutoridad.objects.filter(registro_padron__eleccion_claustro_departamento__eleccion_claustro=claustro, excluida=False).exists():
            redistribuir_autoridades(claustro, usuario)


@transaction.atomic
def responder_asignacion(asignacion, aceptar):
    asignacion.estado = AsignacionAutoridad.Estado.CONFIRMADA if aceptar else AsignacionAutoridad.Estado.RECHAZADA
    asignacion.respondida_en = timezone.now()
    asignacion.save(update_fields=("estado", "respondida_en"))
    if not aceptar and asignacion.candidatura_id:
        asignacion.candidatura.excluida = True
        asignacion.candidatura.save(update_fields=("excluida",))


# La importación solo identifica electores; los datos personales siempre salen del padrón.
def validar_csv_autoridades(contenido, eleccion_claustro, nombre_archivo=""):
    try:
        filas = leer_filas_autoridades(contenido, nombre_archivo)
    except ValueError as error:
        return [], [(None, str(error))]
    if not filas:
        return [], [(None, "El archivo no contiene filas de autoridades.")]
    encabezados = {str(c).strip().casefold(): c for c in filas[0]}
    campo_documento = next((encabezados[n] for n in ("documento", "dni", "cuil") if n in encabezados), None)
    campo_tipo = next((encabezados[n] for n in ("tipo documento", "tipo_documento") if n in encabezados), None)
    campo_legajo = encabezados.get("legajo")
    if not campo_documento and not campo_legajo:
        return [], [(None, "El archivo debe incluir Documento (DNI/CUIL) o Legajo.")]
    normalizadas, errores, vistos = [], [], set()
    for numero, fila in enumerate(filas, start=2):
        doc = (fila.get(campo_documento) or "").strip() if campo_documento else ""
        legajo = (fila.get(campo_legajo) or "").strip() if campo_legajo else ""
        tipo = (fila.get(campo_tipo) or "DNI").strip() if campo_tipo else ("CUIL" if "cuil" in encabezados else "DNI")
        clave = (tipo.casefold(), doc, legajo)
        if not doc and not legajo:
            errores.append((numero, "Debe informarse Documento o Legajo."))
            continue
        if clave in vistos:
            errores.append((numero, "El elector está repetido en el archivo."))
        vistos.add(clave)
        consulta = RegistroPadron.objects.filter(
            eleccion=eleccion_claustro.eleccion,
            eleccion_claustro_departamento__eleccion_claustro=eleccion_claustro,
            activo=True,
        )
        if doc:
            consulta = consulta.filter(elector__tipo_documento__iexact=tipo, elector__dni=doc)
        if legajo:
            consulta = consulta.filter(elector__legajo=legajo)
        coincidencias = consulta.count()
        if coincidencias == 0:
            errores.append((numero, "El elector no pertenece al padron activo de este claustro."))
        elif coincidencias > 1:
            errores.append((numero, "Los identificadores coinciden con más de un elector del padrón."))
        normalizadas.append({"documento": doc, "dni": doc, "tipo_documento": tipo, "legajo": legajo})
    return normalizadas, errores


@transaction.atomic
def importar_autoridades(contenido, eleccion_claustro, usuario, nombre_archivo=""):
    if RegistroPadron.objects.filter(eleccion=eleccion_claustro.eleccion, qr_generado_en__isnull=False).exists():
        return 0, [(None, "La nómina no puede cambiarse después de emitir el padrón imprimible.")]
    filas, errores = validar_csv_autoridades(contenido, eleccion_claustro, nombre_archivo)
    if errores:
        return 0, errores
    creadas = 0
    for fila in filas:
        consulta = RegistroPadron.objects.filter(
            eleccion=eleccion_claustro.eleccion,
            eleccion_claustro_departamento__eleccion_claustro=eleccion_claustro,
            activo=True,
        )
        if fila["documento"]:
            consulta = consulta.filter(elector__tipo_documento__iexact=fila["tipo_documento"], elector__dni=fila["documento"])
        if fila["legajo"]:
            consulta = consulta.filter(elector__legajo=fila["legajo"])
        registro = consulta.get()
        candidatura, creada = CandidaturaAutoridad.objects.get_or_create(
            registro_padron=registro,
            defaults={"cargada_por": usuario},
        )
        if candidatura.excluida:
            candidatura.excluida = False
            candidatura.save(update_fields=("excluida",))
        creadas += int(creada)
    if Mesa.objects.filter(eleccion_claustro_departamento__eleccion_claustro=eleccion_claustro).exists():
        redistribuir_autoridades(eleccion_claustro, usuario)
    return creadas, []
